"""Read the supplied bilingual question tables; never edit the source document."""
import argparse
import hashlib
import json
import re
import sys
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
parser = argparse.ArgumentParser()
parser.add_argument("source")
parser.add_argument("--topic", type=int)
args = parser.parse_args()
doc = Document(args.source)
keys = [dict((int(n), ord(a) - 65) for n, a in re.findall(r"(\d+)\.\s*([ABCD])", p.text))
        for p in doc.paragraphs if len(re.findall(r"(\d+)\.\s*([ABCD])", p.text)) == 10]
titles = []
last = ""
for child in doc.element.body:
    if child.tag.endswith("}p"):
        text = "".join(child.xpath(".//w:t/text()"))
        if text.strip():
            last = text.strip()
    elif child.tag.endswith("}tbl"):
        titles.append(last)
if len(doc.tables) != 8 or len(keys) != 8 or len(titles) != 8:
    raise ValueError("Expected eight source topics and eight complete answer keys.")
rows = []
for ti, table in enumerate(doc.tables):
    if len(table.rows) != 11 or [c.text.strip() for c in table.rows[0].cells] != ["No.", "Conceptual Question", "A", "B", "C", "D"]:
        raise ValueError("Unexpected question table format.")
    for ri, row in enumerate(table.rows[1:], 1):
        cells = [c.text.strip().replace("\r\n", "\n") for c in row.cells]
        marked = [oi for oi, cell in enumerate(row.cells[2:])
                  if cell._tc.xpath('.//w:highlight[@w:val="yellow"]') or cell._tc.xpath('.//w:shd[@w:fill="FFFF00"]')]
        if cells[0] != str(ri) or marked != [keys[ti][ri]]:
            raise ValueError(f"Answer key and highlighted option disagree at topic {ti + 1}, question {ri}.")
        if not all("\n" in value and re.search(r"[\u0600-\u06ff]", value) for value in cells[1:]):
            raise ValueError("A bilingual question or option is incomplete.")
        if len(set(cells[2:])) != 4:
            raise ValueError("Question options must be distinct.")
        rows.append({"id": f"dit-cn-ch1-t{ti + 1:02d}-q{ri:02d}", "topic": titles[ti],
                     "sourceNumber": ri, "question": cells[1], "options": cells[2:], "correctIndex": keys[ti][ri]})
if len(rows) != 80 or len({row["question"] for row in rows}) != 80:
    raise ValueError("Expected 80 unique questions.")
if args.topic is not None:
    if args.topic not in range(8):
        raise ValueError("Topic must be between zero and seven.")
    rows = rows[args.topic * 10:(args.topic + 1) * 10]
with open(args.source, "rb") as source:
    digest = hashlib.file_digest(source, "sha256").hexdigest()
print(json.dumps({"sourceSha256": digest, "chapterNumber": 1, "chapterTitle": "Basic Terminologies", "questions": rows}, ensure_ascii=False))
